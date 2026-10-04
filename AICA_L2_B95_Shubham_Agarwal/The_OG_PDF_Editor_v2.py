"""The OG PDF Editor. Install: py -m pip install PyMuPDF Pillow python-docx openpyxl"""
import os
import re
import tempfile
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path

try:
    import fitz  # PyMuPDF
    from PIL import Image, ImageTk
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    raise SystemExit('Install dependencies first: py -m pip install PyMuPDF Pillow python-docx openpyxl')

BG = '#10192d'
PANEL = '#1b2943'
ACCENT = '#29c7b4'
TEXT = '#f2f7ff'
MUTED = '#b8c9dd'


def page_groups(spec, total):
    """Parse comma-separated single pages and inclusive ranges, preserving groups."""
    groups = []
    for part in spec.split(','):
        part = part.strip()
        match = re.fullmatch(r'(\d+)(?:\s*-\s*(\d+))?', part)
        if not match:
            raise ValueError('Enter pages like 1, 3-5, 8-10.')
        first = int(match.group(1))
        last = int(match.group(2) or first)
        if first < 1 or last > total or first > last:
            raise ValueError(f'Pages must be between 1 and {total}; ranges must ascend.')
        groups.append(list(range(first - 1, last)))
    return groups


def open_pdf(path):
    doc = fitz.open(path)
    if doc.needs_pass:
        doc.close()
        raise ValueError('Password-protected PDFs are not supported.')
    if not doc.is_pdf or not doc.page_count:
        doc.close()
        raise ValueError('Select a nonempty PDF file.')
    return doc


def write_pdf(doc, path, **options):
    """Save via a temporary file, so failed operations do not replace existing files."""
    folder = os.path.dirname(os.path.abspath(path))
    fd, temp = tempfile.mkstemp(prefix='.og_pdf_', suffix='.pdf', dir=folder)
    os.close(fd)
    try:
        doc.save(temp, garbage=4, deflate=True, **options)
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.remove(temp)


class RedactionWindow(tk.Toplevel):
    def __init__(self, master, path):
        super().__init__(master)
        self.title('Redact PDF | The OG PDF Editor')
        self.configure(bg=BG)
        self.geometry('1050x820')
        self.doc = open_pdf(path)
        self.path = path
        self.index = 0
        self.boxes = {}  # page index -> list of PDF-coordinate rectangles
        self.photo = None
        self.scale = 1.0
        self.start = None
        self.preview_rectangle = None
        self.all_pages = tk.BooleanVar(value=False)
        self.closed = False
        self.protocol('WM_DELETE_WINDOW', self.close)
        bar = tk.Frame(self, bg=PANEL, padx=14, pady=12)
        bar.pack(fill='x')
        tk.Label(bar, text='Drag to mark areas for permanent redaction', bg=PANEL,
                 fg=TEXT, font=('Segoe UI', 13, 'bold')).pack(side='left')
        self._button(bar, '◀ Previous', lambda: self.navigate(-1)).pack(side='left', padx=(24, 4))
        self.page_label = tk.Label(bar, bg=PANEL, fg=TEXT)
        self.page_label.pack(side='left', padx=7)
        self._button(bar, 'Next ▶', lambda: self.navigate(1)).pack(side='left')
        tk.Checkbutton(bar, text='Apply new rectangle to all pages', variable=self.all_pages,
                       bg=PANEL, fg=TEXT, selectcolor=BG, activebackground=PANEL,
                       activeforeground=TEXT).pack(side='right')
        tools = tk.Frame(self, bg=BG, padx=14, pady=8)
        tools.pack(fill='x')
        self._button(tools, 'Undo last on this page', self.undo).pack(side='left', padx=4)
        self._button(tools, 'Clear all marks', self.clear).pack(side='left', padx=4)
        self._button(tools, 'Save redacted PDF', self.save).pack(side='right', padx=4)
        self.status = tk.Label(self, text='', bg=BG, fg=MUTED)
        self.status.pack(fill='x')
        holder = tk.Frame(self, bg=BG)
        holder.pack(fill='both', expand=True)
        self.canvas = tk.Canvas(holder, bg='#49566b', cursor='crosshair')
        scroll = ttk.Scrollbar(holder, orient='vertical', command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=scroll.set)
        scroll.pack(side='right', fill='y')
        self.canvas.pack(side='left', fill='both', expand=True)
        self.canvas.bind('<ButtonPress-1>', self.mouse_down)
        self.canvas.bind('<B1-Motion>', self.mouse_move)
        self.canvas.bind('<ButtonRelease-1>', self.mouse_up)
        self.canvas.bind('<MouseWheel>', lambda e: self.canvas.yview_scroll(-int(e.delta / 120), 'units'))
        self.after(120, self.render)

    def _button(self, parent, label, command):
        return tk.Button(parent, text=label, command=command, bg=ACCENT,
                         fg='#10233b', font=('Segoe UI', 10, 'bold'), relief='flat', padx=10, pady=5)

    def render(self):
        page = self.doc[self.index]
        available = max(550, self.canvas.winfo_width() - 42)
        self.scale = min(1.8, available / page.rect.width)
        pix = page.get_pixmap(matrix=fitz.Matrix(self.scale, self.scale), alpha=False)
        self.photo = ImageTk.PhotoImage(Image.frombytes('RGB', (pix.width, pix.height), pix.samples))
        self.canvas.delete('all')
        self.canvas.create_image(16, 16, image=self.photo, anchor='nw')
        self.canvas.configure(scrollregion=(0, 0, pix.width + 32, pix.height + 32))
        for rect in self.boxes.get(self.index, []):
            self.canvas.create_rectangle(16 + rect.x0 * self.scale, 16 + rect.y0 * self.scale,
                                         16 + rect.x1 * self.scale, 16 + rect.y1 * self.scale,
                                         fill='#d94360', stipple='gray25', outline='#ee294c', width=2)
        self.page_label.configure(text=f'Page {self.index + 1} / {self.doc.page_count}')
        count = sum(map(len, self.boxes.values()))
        self.status.configure(text=f'{count} marked area(s). Drag over a page; scrolling is supported.')

    def navigate(self, amount):
        self.index = max(0, min(self.doc.page_count - 1, self.index + amount))
        self.render()

    def coords(self, event):
        return ((self.canvas.canvasx(event.x) - 16) / self.scale,
                (self.canvas.canvasy(event.y) - 16) / self.scale)

    def mouse_down(self, event):
        self.start = self.coords(event)

    def mouse_move(self, event):
        if self.start is None:
            return
        x0, y0 = self.start
        x1, y1 = self.coords(event)
        if self.preview_rectangle:
            self.canvas.delete(self.preview_rectangle)
        self.preview_rectangle = self.canvas.create_rectangle(16 + x0*self.scale, 16 + y0*self.scale,
            16 + x1*self.scale, 16 + y1*self.scale, outline='#f9d66c', width=3)

    def mouse_up(self, event):
        if self.start is None:
            return
        x0, y0 = self.start
        x1, y1 = self.coords(event)
        self.start = None
        source = self.doc[self.index].rect
        rect = fitz.Rect(max(0, min(x0, x1)), max(0, min(y0, y1)),
                         min(source.width, max(x0, x1)), min(source.height, max(y0, y1)))
        if rect.width < 2 or rect.height < 2:
            self.render()
            return
        if self.all_pages.get():
            # Same proportional area, suitable for pages with different dimensions.
            for n in range(self.doc.page_count):
                target = self.doc[n].rect
                scaled = fitz.Rect(rect.x0 * target.width / source.width,
                                   rect.y0 * target.height / source.height,
                                   rect.x1 * target.width / source.width,
                                   rect.y1 * target.height / source.height)
                self.boxes.setdefault(n, []).append(scaled)
        else:
            self.boxes.setdefault(self.index, []).append(rect)
        self.render()

    def undo(self):
        if self.boxes.get(self.index):
            self.boxes[self.index].pop()
            self.render()

    def clear(self):
        self.boxes.clear()
        self.render()

    def save(self):
        if not any(self.boxes.values()):
            messagebox.showwarning('No marks', 'Mark at least one area first.', parent=self)
            return
        output = filedialog.asksaveasfilename(parent=self, title='Save redacted PDF',
            defaultextension='.pdf', filetypes=[('PDF files', '*.pdf')],
            initialfile=f'{Path(self.path).stem}_redacted.pdf')
        if not output:
            return
        if os.path.abspath(output) == os.path.abspath(self.path):
            messagebox.showerror('Choose another file', 'Save to a new filename to protect the original.', parent=self)
            return
        try:
            for n, rects in self.boxes.items():
                page = self.doc[n]
                for rect in rects:
                    page.add_redact_annot(rect, fill=(1, 1, 1), cross_out=False)
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,
                                      graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                                      text=fitz.PDF_REDACT_TEXT_REMOVE)
            write_pdf(self.doc, output, clean=True)
            messagebox.showinfo('Saved', f'Redacted PDF saved:\n{output}', parent=self)
            self.close()
        except Exception as exc:
            messagebox.showerror('Save failed', str(exc), parent=self)

    def close(self):
        if not self.closed:
            self.closed = True
            self.doc.close()
        self.destroy()


class PDFEditor(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('The OG PDF Editor')
        self.geometry('790x670')
        self.minsize(710, 610)
        self.configure(bg=BG)
        style = ttk.Style(self)
        style.theme_use('clam')
        style.configure('Accent.TCombobox', fieldbackground='white', background=ACCENT)
        header = tk.Frame(self, bg='#243c66', padx=30, pady=20)
        header.pack(fill='x')
        tk.Label(header, text='The OG PDF Editor', bg='#243c66', fg=TEXT,
                 font=('Segoe UI', 24, 'bold')).pack(anchor='w')
        tk.Label(header, text='Merge · Split · Extract · Rotate · Redact · Compress',
                 bg='#243c66', fg='#a8f1e5', font=('Segoe UI', 11)).pack(anchor='w')
        body = tk.Frame(self, bg=BG, padx=24, pady=16)
        body.pack(fill='both', expand=True)
        self.tab = ttk.Notebook(body)
        self.tab.pack(fill='both', expand=True)
        self.setup_tab('Merge', 'Choose PDFs in the intended order. Move selected files up or down.', self.merge_ui)
        self.setup_tab('Split', 'Each comma-separated page or range becomes its own PDF.', self.split_ui)
        self.setup_tab('Extract', 'Combine selected pages into one PDF in the order entered.', self.extract_ui)
        self.setup_tab('Rotate', 'Rotate selected pages clockwise, preserving other pages.', self.rotate_ui)
        self.setup_tab('Redact', 'Mark rectangles in the preview, then permanently remove underlying content.', self.redact_ui)
        self.setup_tab('Compress', 'Try lossless optimization for a selected file size.', self.compress_ui)
        self.setup_tab('Convert', 'Export PDF pages to images, Word, Excel or plain text.', self.convert_ui)
        tk.Label(self, text='Files stay on your computer  •  Originals are never overwritten',
                 bg=BG, fg=MUTED, font=('Segoe UI', 9)).pack(pady=(0, 12))

    def setup_tab(self, name, description, builder):
        frame = tk.Frame(self.tab, bg=PANEL, padx=24, pady=20)
        self.tab.add(frame, text=f'  {name}  ')
        tk.Label(frame, text=name + ' PDF' if name != 'Redact' else 'Redact PDF',
                 fg=TEXT, bg=PANEL, font=('Segoe UI', 17, 'bold')).pack(anchor='w')
        tk.Label(frame, text=description, fg=MUTED, bg=PANEL,
                 wraplength=610, justify='left').pack(anchor='w', pady=(5, 18))
        builder(frame)

    def label(self, frame, value):
        tk.Label(frame, text=value, bg=PANEL, fg=TEXT, font=('Segoe UI', 10, 'bold')).pack(anchor='w', pady=(12, 4))

    def button(self, frame, value, command):
        button = tk.Button(frame, text=value, command=command, bg=ACCENT,
                           fg='#11283e', relief='flat', padx=16, pady=9,
                           font=('Segoe UI', 10, 'bold'), cursor='hand2')
        button.pack(anchor='w', pady=10)
        return button

    def entry(self, frame, placeholder):
        var = tk.StringVar(value=placeholder)
        tk.Entry(frame, textvariable=var, width=48, font=('Segoe UI', 11),
                 bg='white', fg='#192b43', relief='flat').pack(anchor='w', ipady=8)
        return var

    def choose_one(self, label):
        path = filedialog.askopenfilename(parent=self, title=label, filetypes=[('PDF files', '*.pdf')])
        if not path:
            return None
        try:
            with open_pdf(path) as doc:
                count = doc.page_count
        except Exception as exc:
            messagebox.showerror('Cannot open PDF', str(exc), parent=self)
            return None
        return path, count

    def source(self, frame):
        var = tk.StringVar(value='No PDF selected')
        self.button(frame, 'Choose PDF', lambda: self.select_source(var))
        tk.Label(frame, textvariable=var, bg=PANEL, fg=MUTED,
                 wraplength=600, justify='left').pack(anchor='w')
        return var

    def select_source(self, var):
        selected = self.choose_one('Choose a PDF')
        if selected:
            var.set(f'{selected[0]}  ({selected[1]} pages)')
            var.path, var.count = selected

    def selected(self, var):
        if not hasattr(var, 'path'):
            raise ValueError('Choose a PDF first.')
        return var.path, var.count

    def target(self, source, suffix):
        return filedialog.asksaveasfilename(parent=self, title='Save output',
            initialfile=Path(source).stem + suffix + '.pdf',
            defaultextension='.pdf', filetypes=[('PDF files', '*.pdf')])

    def execute(self, action):
        try:
            result = action()
            if result:
                messagebox.showinfo('Complete', result, parent=self)
        except Exception as exc:
            messagebox.showerror('Unable to complete', str(exc), parent=self)

    def merge_ui(self, frame):
        items = tk.Listbox(frame, height=10, width=78, selectmode='browse',
                           bg='#f5f8fc', fg='#1b2943', font=('Segoe UI', 10))
        items.pack(fill='x')
        controls = tk.Frame(frame, bg=PANEL)
        controls.pack(fill='x')
        for caption, command in [('Add PDFs', lambda: [items.insert('end', x) for x in
                filedialog.askopenfilenames(parent=self, filetypes=[('PDF files', '*.pdf')])]),
                ('Remove', lambda: items.delete(tk.ACTIVE) if items.curselection() else None),
                ('Move up', lambda: self.move(items, -1)),
                ('Move down', lambda: self.move(items, 1))]:
            tk.Button(controls, text=caption, command=command, bg='#354f76', fg=TEXT,
                      relief='flat', padx=9, pady=5).pack(side='left', padx=3, pady=8)
        self.button(frame, 'Merge and save', lambda: self.execute(lambda: self.merge(items)))

    @staticmethod
    def move(items, step):
        if not items.curselection():
            return
        old = items.curselection()[0]
        new = old + step
        if 0 <= new < items.size():
            value = items.get(old)
            items.delete(old)
            items.insert(new, value)
            items.selection_set(new)

    def merge(self, items):
        paths = items.get(0, 'end')
        if len(paths) < 2:
            raise ValueError('Choose at least two PDFs.')
        output = self.target(paths[0], '_merged')
        if not output:
            return None
        if os.path.abspath(output) in {os.path.abspath(p) for p in paths}:
            raise ValueError('Choose a new filename for the merged PDF.')
        result = fitz.open()
        try:
            for path in paths:
                with open_pdf(path) as source:
                    result.insert_pdf(source)
            write_pdf(result, output)
        finally:
            result.close()
        return f'Merged PDF saved:\n{output}'

    def split_ui(self, frame):
        source = self.source(frame)
        self.label(frame, 'Page groups (example: 1, 3-5, 8-10)')
        spec = self.entry(frame, '1, 3-5, 8-10')
        self.button(frame, 'Split into PDFs', lambda: self.execute(lambda: self.split(source, spec.get())))

    def split(self, source_var, spec):
        path, count = self.selected(source_var)
        groups = page_groups(spec, count)
        folder = filedialog.askdirectory(parent=self, title='Select output folder')
        if not folder:
            return None
        with open_pdf(path) as original:
            # Validate names before writing any output.
            names = [os.path.join(folder, f'{Path(path).stem}_part_{i:02d}.pdf') for i in range(1, len(groups)+1)]
            if any(os.path.abspath(name) == os.path.abspath(path) for name in names):
                raise ValueError('Choose another output folder.')
            if any(os.path.exists(name) for name in names) and not messagebox.askyesno(
                    'Replace files?', 'Some output files already exist. Replace them?', parent=self):
                return None
            for pages, name in zip(groups, names):
                output = fitz.open()
                try:
                    for p in pages:
                        output.insert_pdf(original, from_page=p, to_page=p)
                    write_pdf(output, name)
                finally:
                    output.close()
        return f'Created {len(groups)} PDF files in:\n{folder}'

    def extract_ui(self, frame):
        source = self.source(frame)
        self.label(frame, 'Pages to extract (example: 1, 3-5, 8-10)')
        spec = self.entry(frame, '1, 3-5, 8-10')
        self.button(frame, 'Extract and save', lambda: self.execute(lambda: self.extract(source, spec.get())))

    def extract(self, source_var, spec):
        path, count = self.selected(source_var)
        pages = [p for group in page_groups(spec, count) for p in group]
        output_path = self.target(path, '_extracted')
        if not output_path:
            return None
        if os.path.abspath(output_path) == os.path.abspath(path):
            raise ValueError('Choose a different output filename.')
        with open_pdf(path) as original:
            output = fitz.open()
            try:
                for p in pages:
                    output.insert_pdf(original, from_page=p, to_page=p)
                write_pdf(output, output_path)
            finally:
                output.close()
        return f'Extracted {len(pages)} pages to:\n{output_path}'

    def rotate_ui(self, frame):
        source = self.source(frame)
        self.label(frame, 'Pages to rotate (example: 2, 4-6)')
        spec = self.entry(frame, '2, 4-6')
        self.label(frame, 'Clockwise angle')
        angle = tk.StringVar(value='90°')
        ttk.Combobox(frame, textvariable=angle, state='readonly', width=15,
                     values=('90°', '180°', '270°')).pack(anchor='w')
        self.button(frame, 'Rotate and save', lambda: self.execute(
            lambda: self.rotate(source, spec.get(), angle.get())))

    def rotate(self, source_var, spec, angle):
        path, count = self.selected(source_var)
        pages = set(p for group in page_groups(spec, count) for p in group)
        output = self.target(path, '_rotated')
        if not output:
            return None
        if os.path.abspath(output) == os.path.abspath(path):
            raise ValueError('Choose a different output filename.')
        degrees = int(angle.rstrip('°'))
        with open_pdf(path) as doc:
            for index in pages:
                page = doc[index]
                page.set_rotation((page.rotation + degrees) % 360)
            write_pdf(doc, output)
        return f'Rotated {len(pages)} pages:\n{output}'

    def redact_ui(self, frame):
        self.button(frame, 'Choose PDF and open preview', lambda: self.execute(self.redact))
        tk.Label(frame, text='Draw one or more rectangles. Turn on “all pages” before drawing\n'
                 'to place the same proportional area on every page.', bg=PANEL,
                 fg=MUTED, justify='left').pack(anchor='w')

    def redact(self):
        choice = self.choose_one('Choose PDF to redact')
        if choice:
            RedactionWindow(self, choice[0])
        return None

    def compress_ui(self, frame):
        source = self.source(frame)
        self.label(frame, 'Target maximum size')
        target = tk.StringVar(value='500 KB')
        ttk.Combobox(frame, textvariable=target, state='readonly', width=16,
            values=tuple(f'{size} KB' for size in range(500, 10001, 500))).pack(anchor='w')
        self.button(frame, 'Optimize and save', lambda: self.execute(
            lambda: self.compress(source, target.get())))
        tk.Label(frame, text='Lossless compression keeps page clarity. Some PDFs cannot reach a chosen size\n'
                 'without lowering image quality; the app reports the actual size.',
                 bg=PANEL, fg=MUTED, justify='left').pack(anchor='w', pady=12)

    def compress(self, source_var, target):
        path, _ = self.selected(source_var)
        output = self.target(path, '_compressed')
        if not output:
            return None
        if os.path.abspath(output) == os.path.abspath(path):
            raise ValueError('Choose a different output filename.')
        limit = int(target.split()[0]) * 1024
        # Save multiple lossless candidates to temporary files, retain only the smallest.
        candidates = []
        try:
            with open_pdf(path) as doc:
                for extra in ({}, {'deflate_images': True, 'deflate_fonts': True, 'use_objstms': 1}):
                    fd, temp = tempfile.mkstemp(prefix='.og_candidate_', suffix='.pdf',
                                                dir=os.path.dirname(os.path.abspath(output)))
                    os.close(fd)
                    candidates.append(temp)
                    doc.save(temp, garbage=4, deflate=True, **extra)
            best = min(candidates, key=os.path.getsize)
            size = os.path.getsize(best)
            original_size = os.path.getsize(path)
            if size >= original_size:
                return (f'The original is already smaller ({original_size / 1024:.0f} KB) '
                        'than the lossless optimized version. No output was saved.')
            os.replace(best, output)
            note = 'Target reached.' if size <= limit else ('The selected size could not be reached '
                'without reducing image quality. The smallest lossless result was saved.')
            return f'{note}\nOriginal: {original_size / 1024:.0f} KB\nResult: {size / 1024:.0f} KB\n{output}'
        finally:
            for temp in candidates:
                if os.path.exists(temp):
                    os.remove(temp)

    def convert_ui(self, frame):
        source = self.source(frame)
        self.label(frame, 'Output format')
        format_var = tk.StringVar(value='JPG images')
        ttk.Combobox(frame, textvariable=format_var, state='readonly', width=22,
                     values=('JPG images', 'PNG images', 'Word (.docx)',
                             'Excel (.xlsx)', 'Plain text (.txt)')).pack(anchor='w')
        self.button(frame, 'Convert PDF', lambda: self.execute(
            lambda: self.convert(source, format_var.get())))
        tk.Label(frame, text='JPG / PNG: one image per page, rendered at 200 DPI.\n'
                 'Word: editable extracted text, with page breaks. Excel: detected tables\n'
                 'on separate sheets; pages without tables are saved as text rows.\n'
                 'Scanned pages without selectable text need OCR, which is not included.\n'
                 'Complex layouts may not carry over to Word or Excel.',
                 bg=PANEL, fg=MUTED, justify='left').pack(anchor='w', pady=12)

    def convert(self, source_var, output_format):
        path, _ = self.selected(source_var)
        if output_format in ('JPG images', 'PNG images'):
            folder = filedialog.askdirectory(parent=self, title='Choose folder for page images')
            if not folder:
                return None
            extension = '.jpg' if output_format == 'JPG images' else '.png'
            with open_pdf(path) as pdf:
                destinations = [os.path.join(folder, f'{Path(path).stem}_page_{i + 1:03d}{extension}')
                                for i in range(pdf.page_count)]
                if any(os.path.exists(name) for name in destinations) and not messagebox.askyesno(
                        'Replace images?', 'Some page images already exist. Replace them?', parent=self):
                    return None
                for page, destination in zip(pdf, destinations):
                    pix = page.get_pixmap(matrix=fitz.Matrix(200 / 72, 200 / 72), alpha=False)
                    if extension == '.png':
                        pix.save(destination)
                    else:
                        image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
                        image.save(destination, 'JPEG', quality=92, optimize=True)
            return f'Saved {len(destinations)} page images in:\n{folder}'

        suffix = {'Word (.docx)': '.docx', 'Excel (.xlsx)': '.xlsx',
                  'Plain text (.txt)': '.txt'}[output_format]
        output = filedialog.asksaveasfilename(parent=self, title='Save converted file',
            initialfile=Path(path).stem + '_converted' + suffix,
            defaultextension=suffix, filetypes=[(output_format, '*' + suffix)])
        if not output:
            return None
        if os.path.abspath(output) == os.path.abspath(path):
            raise ValueError('Choose a new output filename.')
        fd, temporary = tempfile.mkstemp(prefix='.og_conversion_', suffix=suffix,
                                         dir=os.path.dirname(os.path.abspath(output)))
        os.close(fd)
        try:
            with open_pdf(path) as pdf:
                if suffix == '.docx':
                    document = Document()
                    for number, page in enumerate(pdf):
                        if number:
                            document.add_page_break()
                        document.add_heading(f'Page {number + 1}', level=2)
                        for block in page.get_text('blocks', sort=True):
                            if len(block) > 4 and block[4].strip():
                                document.add_paragraph(block[4].strip())
                    document.save(temporary)
                elif suffix == '.txt':
                    with open(temporary, 'w', encoding='utf-8') as output_file:
                        for number, page in enumerate(pdf):
                            output_file.write(f'\n--- Page {number + 1} ---\n')
                            output_file.write(page.get_text(sort=True))
                            output_file.write('\n')
                else:
                    workbook = Workbook()
                    workbook.remove(workbook.active)
                    for number, page in enumerate(pdf):
                        sheet = workbook.create_sheet(f'Page {number + 1}')
                        tables = page.find_tables().tables
                        if tables:
                            for index, table in enumerate(tables):
                                sheet.append([f'Table {index + 1}'])
                                sheet.cell(sheet.max_row, 1).font = Font(bold=True, color='FFFFFF')
                                sheet.cell(sheet.max_row, 1).fill = PatternFill('solid', fgColor='243C66')
                                for row in table.extract():
                                    sheet.append([str(value) if value is not None else '' for value in row])
                                sheet.append([])
                        else:
                            sheet.append(['Extracted text (no table detected)'])
                            for line in page.get_text(sort=True).splitlines():
                                sheet.append([line])
                        sheet.column_dimensions['A'].width = 48
                        for column in sheet.columns:
                            if column[0].column > 1:
                                sheet.column_dimensions[column[0].column_letter].width = 24
                        sheet.freeze_panes = 'A2'
                    workbook.save(temporary)
            os.replace(temporary, output)
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)
        return f'Converted file saved:\n{output}'


if __name__ == '__main__':
    PDFEditor().mainloop()
